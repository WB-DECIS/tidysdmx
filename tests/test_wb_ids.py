"""Tests for the WB_IDS mock dataset harness (benchmarks/wb_ids/, not shipped).

The fixtures live here rather than in ``tests/fixtures``: they import
``benchmarks``, which the sdist does not carry, and conftest.py must import
cleanly from an unpacked sdist. This module is excluded from the sdist for the
same reason as ``test_benchmarks.py``.
"""

import csv
import json
import re
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from pysdmx.model import Role

from benchmarks.wb_ids import (
    PROFILE_JSON,
    SAMPLE_CSV,
    SAMPLE_ROWS,
    SCHEMA_JSON,
    TEMPLATE_XLSX,
)
from benchmarks.wb_ids.generate import _codes, _universe, generate, load_profile
from benchmarks.wb_ids.memory import PeakRss
from benchmarks.wb_ids.pipeline import (
    describe_mode,
    main,
    read_csv_kwargs,
    requested_series,
    run_pipeline,
    to_markdown,
)
from benchmarks.wb_ids.profile import COLUMNS, dumps, profile_csv
from benchmarks.wb_ids.schema import (
    from_spec,
    load_schema,
    reachable_codes,
    to_spec,
)
from benchmarks.wb_ids.template import scrub
from tidysdmx import parse_mapping_template_wb

KEY = ["Country_Code", "Counterpart_Area_Code", "Series", "Year"]
FIDELITY_ROWS = 50_000


@pytest.fixture(scope="module")
def real_profile():
    return load_profile(PROFILE_JSON)


@pytest.fixture(scope="module")
def mock_csv(tmp_path_factory):
    # Generating builds the full-size series universe (a few seconds), so the
    # file is made once per module; tests only read it.
    path = tmp_path_factory.mktemp("wb_ids") / "mock.csv"
    generate(path, FIDELITY_ROWS, seed=0)
    return path


@pytest.fixture(scope="module")
def mock_profile(mock_csv):
    return profile_csv(mock_csv)


@pytest.fixture(scope="module")
def pipeline_result():
    return run_pipeline(SAMPLE_CSV)


@pytest.fixture(scope="module")
def outputs(tmp_path_factory):
    """Each mode's run and standardised output file, on the committed sample."""
    folder = tmp_path_factory.mktemp("wb_ids_modes")
    runs = {
        "whole": {},
        "chunked": {"chunk_rows": 1000},
        "chunked-lean": {"chunk_rows": 1000, "skip_unused": True, "categories": True},
        "whole-lean": {"skip_unused": True, "categories": True},
    }
    results = {}
    for name, options in runs.items():
        out_csv = folder / f"{name}.csv"
        results[name] = (run_pipeline(SAMPLE_CSV, out_csv=out_csv, **options), out_csv)
    return results


def _read(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, dtype=str, keep_default_na=False)


def _write_csv(path: Path, rows: list[list[str]]) -> Path:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, quoting=csv.QUOTE_ALL, lineterminator="\n")
        writer.writerow(COLUMNS)
        writer.writerows(rows)
    return path


class TestGenerate:
    def test_generate_writes_exactly_the_rows_asked_for(self, mock_csv):
        assert len(_read(mock_csv)) == FIDELITY_ROWS

    def test_generate_reproduces_the_committed_sample(self, tmp_path):
        path = tmp_path / "sample.csv"
        generate(path, SAMPLE_ROWS, seed=0)
        assert path.read_bytes() == SAMPLE_CSV.read_bytes()

    def test_generate_differs_by_seed(self, tmp_path):
        generate(tmp_path / "a.csv", 2_000, seed=1)
        generate(tmp_path / "b.csv", 2_000, seed=2)
        assert (tmp_path / "a.csv").read_bytes() != (tmp_path / "b.csv").read_bytes()

    def test_generate_keys_are_unique(self, mock_csv):
        assert not _read(mock_csv).duplicated(KEY).any()

    def test_generate_header_and_quoting_match_the_real_file(self, mock_csv):
        header = mock_csv.read_bytes().split(b"\n", 1)[0]
        assert header == b",".join(b'"%s"' % name.encode() for name in COLUMNS)

    def test_generate_writes_lf_line_endings(self, mock_profile):
        assert mock_profile["file"]["line_terminator"] == "LF"

    def test_generate_quotes_every_field(self, mock_profile):
        assert mock_profile["file"]["quoted"]

    def test_generate_rejects_zero_rows(self, tmp_path):
        with pytest.raises(ValueError, match="rows must be at least 1"):
            generate(tmp_path / "x.csv", 0)

    def test_generate_rejects_more_rows_than_the_universe_holds(
        self, tmp_path, real_profile
    ):
        # Twice every real series in every year: more than any universe holds.
        too_many = (
            2 * len(real_profile["years"]) * real_profile["keys"]["distinct_triples"]
        )
        with pytest.raises(ValueError, match="exceeds"):
            generate(tmp_path / "x.csv", too_many)


class TestMockFidelity:
    """A mock's profile resembles the real one, scaled to its size."""

    def test_mock_uses_only_real_codes(self, mock_profile, real_profile):
        for section in ("countries", "counterparts", "series"):
            assert set(mock_profile[section]) <= set(real_profile[section])

    def test_mock_series_names_match_the_real_ones(self, mock_profile, real_profile):
        for code, entry in mock_profile["series"].items():
            assert entry["name"] == real_profile["series"][code]["name"]

    def test_mock_covers_every_real_year(self, mock_profile, real_profile):
        assert set(mock_profile["years"]) == set(real_profile["years"])

    def test_mock_year_shares_track_the_real_ones(self, mock_profile, real_profile):
        real_total = real_profile["file"]["rows"]
        distance = sum(
            abs(mock_profile["years"].get(year, 0) / FIDELITY_ROWS - n / real_total)
            for year, n in real_profile["years"].items()
        )
        assert distance / 2 < 0.1

    def test_mock_bytes_per_row_track_the_real_file(self, mock_profile, real_profile):
        real = real_profile["file"]["bytes_per_row"]
        assert mock_profile["file"]["bytes_per_row"] == pytest.approx(real, rel=0.1)

    def test_mock_zero_share_tracks_the_real_file(self, mock_profile, real_profile):
        real = real_profile["values"]["zero"] / real_profile["file"]["rows"]
        mock = mock_profile["values"]["zero"] / FIDELITY_ROWS
        assert mock == pytest.approx(real, abs=0.02)

    def test_universe_gives_the_largest_counterpart_the_most_series(self, real_profile):
        # Checked on the series universe rather than on a sample: a 50k-row
        # mock holds only ~1,200 series, too few to rank 304 counterparts.
        codes = _codes(real_profile)
        _, counterpart, _ = _universe(real_profile, codes, np.random.default_rng(0))
        real = real_profile["counterparts"]
        largest = codes.counterparts[int(np.bincount(counterpart).argmax())]
        assert largest == max(real, key=real.get)

    def test_mock_has_no_duplicate_keys(self, mock_profile):
        assert mock_profile["keys"]["duplicates"] == 0

    def test_mock_partitions_hold_real_partitions_countries(
        self, mock_profile, real_profile
    ):
        real_runs = [set(run["countries"]) for run in real_profile["sort"]["runs"]]
        for run in mock_profile["sort"]["runs"]:
            assert any(set(run["countries"]) <= real for real in real_runs)

    def test_mock_is_ordered_within_partitions(self, mock_profile):
        sort = mock_profile["sort"]
        assert sort["country_order_breaks"] == sort["counterpart_order_breaks"] == 0

    def test_mock_values_are_one_decimal_text(self, mock_profile):
        assert mock_profile["values"]["one_decimal"] == FIDELITY_ROWS


class TestProfile:
    def test_profile_counts_duplicate_keys(self, tmp_path):
        row = ["AGO", "WLD", "S1", "Name 1", "1970", "0.0", "", ""]
        path = _write_csv(tmp_path / "dup.csv", [row, row])
        assert profile_csv(path)["keys"]["duplicates"] == 1

    def test_profile_finds_year_ordered_partitions(self, tmp_path):
        rows = [
            ["AGO", "WLD", "S1", "N", "1970", "1.0", "", ""],
            ["AGO", "WLD", "S1", "N", "1971", "1.0", "", ""],
            ["BEN", "WLD", "S1", "N", "1970", "1.0", "", ""],
        ]
        runs = profile_csv(_write_csv(tmp_path / "p.csv", rows))["sort"]["runs"]
        assert [run["countries"] for run in runs] == [["AGO"], ["BEN"]]

    def test_profile_rejects_another_layout(self, tmp_path):
        path = tmp_path / "other.csv"
        path.write_text('"a","b"\n"1","2"\n', encoding="utf-8")
        with pytest.raises(ValueError, match="Expected columns"):
            profile_csv(path)

    def test_dumps_keeps_number_lists_on_one_line(self):
        assert '"counts": [1, 2, 3]' in dumps({"counts": [1, 2, 3]})

    def test_committed_profile_is_aggregates_only(self):
        # Aggregates are bounded by the code vocabularies, not by the row count.
        assert PROFILE_JSON.stat().st_size < 1_000_000


class TestSchema:
    def test_load_schema_has_the_ds_ids_dimensions(self):
        schema = load_schema(SCHEMA_JSON)
        dimensions = [c.id for c in schema.components if c.role == Role.DIMENSION]
        assert dimensions[-1] == "TIME_PERIOD"
        assert {"REF_AREA", "INDICATOR", "COUNTERPART_AREA"} <= set(dimensions)

    def test_load_schema_keeps_every_code(self):
        spec = json.loads(SCHEMA_JSON.read_text(encoding="utf-8"))
        expected = {c["id"]: len(c["codes"]) for c in spec["components"] if c["codes"]}
        schema = load_schema(SCHEMA_JSON)
        loaded = {
            c.id: len(c.enumeration.items)
            for c in schema.components
            if c.enumeration is not None
        }
        assert loaded == expected

    def test_spec_round_trips_except_real_codelist_sizes(self):
        spec = json.loads(SCHEMA_JSON.read_text(encoding="utf-8"))
        rebuilt = to_spec(from_spec(spec), spec["source"])
        # codes_total is the untrimmed size, which the trimmed codes cannot carry.
        for component in (*spec["components"], *rebuilt["components"]):
            component.pop("codes_total")
        assert rebuilt == spec

    def test_to_spec_trims_to_the_kept_codes(self):
        spec = to_spec(load_schema(SCHEMA_JSON), "test", keep={"REF_AREA": {"AGO"}})
        ref_area = next(c for c in spec["components"] if c["id"] == "REF_AREA")
        assert ref_area["codes"] == ["AGO"]

    def test_to_spec_keeps_one_code_when_none_can_occur(self):
        # pysdmx reads an empty codelist as none, which would skip validation.
        spec = to_spec(load_schema(SCHEMA_JSON), "test", keep={"REF_AREA": set()})
        ref_area = next(c for c in spec["components"] if c["id"] == "REF_AREA")
        assert len(ref_area["codes"]) == 1

    def test_reachable_codes_takes_fixed_values(self, real_profile):
        mappings = parse_mapping_template_wb(TEMPLATE_XLSX)
        assert reachable_codes(real_profile, mappings)["FREQ"] == {"A"}

    def test_reachable_codes_takes_implicit_source_codes(self, real_profile):
        mappings = parse_mapping_template_wb(TEMPLATE_XLSX)
        reachable = reachable_codes(real_profile, mappings)["REF_AREA"]
        assert reachable == set(real_profile["countries"])

    def test_reachable_codes_takes_representation_targets(self, real_profile):
        mappings = parse_mapping_template_wb(TEMPLATE_XLSX)
        targets = set(mappings["REP_MAPPING"]["T:INDICATOR"].dropna())
        assert reachable_codes(real_profile, mappings)["INDICATOR"] == targets


class TestScrubTemplate:
    @pytest.fixture
    def real_like(self, tmp_path):
        workbook = Workbook()
        info = workbook.active
        info.title = "INFO"
        info.append(["TTL NAME", "Jane Doe"])
        info.append(["TTL EMAIL", "jdoe@example.com"])
        info["B2"].hyperlink = "mailto:jdoe@example.com"
        comp = workbook.create_sheet("COMP_MAPPING")
        comp.append(["SOURCE", "TARGET", "MAPPING_RULES"])
        comp.append(["Country_code", "REF_AREA", "implicit"])
        rep = workbook.create_sheet("REP_MAPPING")
        rep.append(["S:Series", "T:INDICATOR"])
        rep.append(["DT_X_CD\xa0\xa0", "DT_X"])
        requested = workbook.create_sheet("REQUESTED_IND_RAW")
        requested.append(["Series"])
        requested.append(["DT_X_CD\xa0"])
        workbook.properties.creator = "Jane Doe"
        path = tmp_path / "real.xlsx"
        workbook.save(path)
        out = tmp_path / "mock.xlsx"
        scrub(path, out)
        return out

    def test_scrub_replaces_the_ttl_name_and_email(self, real_like):
        info = load_workbook(real_like)["INFO"]
        assert [info["B1"].value, info["B2"].value] == [
            "Mock Task Team Leader",
            "ttl@example.org",
        ]

    def test_scrub_leaves_no_personal_data_anywhere(self, real_like):
        with zipfile.ZipFile(real_like) as archive:
            text = b"".join(archive.read(name) for name in archive.namelist())
        assert not re.search(rb"Jane|jdoe", text)

    def test_scrub_respells_sources_as_the_csv_columns(self, real_like):
        assert load_workbook(real_like)["COMP_MAPPING"]["A2"].value == "Country_Code"

    def test_scrub_strips_codes(self, real_like):
        workbook = load_workbook(real_like)
        codes = [
            workbook["REP_MAPPING"]["A2"].value,
            workbook["REQUESTED_IND_RAW"]["A2"].value,
        ]
        assert codes == ["DT_X_CD", "DT_X_CD"]

    def test_scrub_rejects_a_workbook_without_comp_mapping(self, tmp_path):
        workbook = Workbook()
        workbook.active.title = "INFO"
        workbook.save(tmp_path / "bad.xlsx")
        with pytest.raises(ValueError, match="no COMP_MAPPING sheet"):
            scrub(tmp_path / "bad.xlsx", tmp_path / "out.xlsx")

    def test_committed_template_has_no_mailto_link(self):
        with zipfile.ZipFile(TEMPLATE_XLSX) as archive:
            text = b"".join(archive.read(name) for name in archive.namelist())
        assert b"mailto:" not in text


class TestPipeline:
    def test_pipeline_runs_every_step(self, pipeline_result):
        assert [step["step"] for step in pipeline_result["steps"]] == [
            "read_template",
            "load",
            "filter",
            "structure_map",
            "map",
            "standardize",
            "validate",
        ]

    def test_pipeline_loads_the_whole_sample(self, pipeline_result):
        assert pipeline_result["steps"][1]["rows"] == SAMPLE_ROWS

    def test_pipeline_keeps_every_filtered_row(self, pipeline_result):
        rows = {step["step"]: step["rows"] for step in pipeline_result["steps"]}
        assert rows["filter"] == rows["map"] == rows["standardize"]

    def test_pipeline_output_has_every_mandatory_column(self, pipeline_result):
        # REF_AREA among them: the real template's `Country_code` would drop it.
        assert "mandatory_columns" not in pipeline_result["validation"]["by_type"]

    def test_pipeline_finds_only_the_real_codelist_gaps(self, pipeline_result):
        assert set(pipeline_result["validation"]["by_type"]) <= {"codelist_ids"}

    def test_to_markdown_has_a_row_per_step(self, pipeline_result):
        table = to_markdown(pipeline_result)
        assert all(f"| {s['step']} |" in table for s in pipeline_result["steps"])

    def test_requested_series_needs_the_sheet(self):
        with pytest.raises(ValueError, match="no REQUESTED_IND_RAW sheet"):
            requested_series({})

    def test_pipeline_records_the_whole_frame_mode(self, pipeline_result):
        assert pipeline_result["input"]["chunk_rows"] is None
        assert pipeline_result["input"]["chunks"] == 1
        assert describe_mode(pipeline_result) == "whole frame"


class TestChunkedPipeline:
    """The chunked run is a drop-in for the whole-frame run."""

    def test_chunked_pipeline_runs_every_step_in_order(self, outputs):
        result, _ = outputs["chunked"]
        assert [step["step"] for step in result["steps"]] == [
            "read_template",
            "load",
            "filter",
            "structure_map",
            "map",
            "standardize",
            "validate",
            "write",
        ]

    def test_chunked_pipeline_reads_the_file_in_chunks(self, outputs):
        result, _ = outputs["chunked"]
        assert result["input"]["chunk_rows"] == 1000
        assert result["input"]["chunks"] == SAMPLE_ROWS // 1000
        assert describe_mode(result) == "chunks of 1,000 rows (4 chunks)"

    def test_chunked_pipeline_sums_rows_over_chunks(self, outputs):
        whole, _ = outputs["whole"]
        chunked, _ = outputs["chunked"]
        for name in ("load", "filter", "map", "standardize"):
            rows = {step["step"]: step["rows"] for step in chunked["steps"]}
            assert rows[name] == next(
                step["rows"] for step in whole["steps"] if step["step"] == name
            )

    @pytest.mark.parametrize("name", ["chunked", "chunked-lean", "whole-lean"])
    def test_pipeline_modes_find_the_same_errors(self, outputs, name):
        whole, _ = outputs["whole"]
        other, _ = outputs[name]
        assert other["validation"] == whole["validation"]

    @pytest.mark.parametrize("name", ["chunked", "chunked-lean", "whole-lean"])
    def test_pipeline_modes_write_identical_output(self, outputs, name):
        _, whole_csv = outputs["whole"]
        _, other_csv = outputs[name]
        assert other_csv.read_bytes() == whole_csv.read_bytes()

    def test_lean_mode_is_described(self, outputs):
        result, _ = outputs["chunked-lean"]
        assert describe_mode(result) == (
            "chunks of 1,000 rows (4 chunks), without Series_Name, "
            "code columns as category"
        )
        assert "chunks of 1,000 rows" in to_markdown(result).splitlines()[0]

    def test_read_csv_kwargs_drop_series_name_and_type_the_codes(self):
        kwargs = read_csv_kwargs(skip_unused=True, categories=True)
        assert "Series_Name" not in kwargs["usecols"]
        assert kwargs["dtype"]["Series"] == "category"
        assert kwargs["dtype"]["footnote"] == "str"
        assert read_csv_kwargs(skip_unused=False, categories=False) == {}

    def test_run_pipeline_rejects_a_non_positive_chunk_size(self):
        with pytest.raises(ValueError, match="chunk_rows must be positive"):
            run_pipeline(SAMPLE_CSV, chunk_rows=0)

    def test_chunked_pipeline_refuses_an_input_with_no_requested_series(self, tmp_path):
        """Fail, as whole-frame mode does, rather than report a clean run.

        An untouched stale output file would otherwise pass for the result.
        """
        rows = [
            ["AFG", "WLD", "NOT_REQUESTED", "Nothing", "2020", "1.0", "", ""],
            ["AFG", "WLD", "NOT_REQUESTED", "Nothing", "2021", "2.0", "", ""],
        ]
        csv = _write_csv(tmp_path / "unrequested.csv", rows)
        stale = tmp_path / "out.csv"
        stale.write_text("stale\n", encoding="utf-8")
        with pytest.raises(ValueError, match="No rows matched the requested series"):
            run_pipeline(csv, chunk_rows=1, out_csv=stale)
        assert stale.read_text(encoding="utf-8") == "stale\n"

    def test_main_takes_the_chunked_flags(self, tmp_path, capsys):
        out = tmp_path / "run.json"
        code = main(
            [
                "--chunk-rows",
                "2000",
                "--skip-unused",
                "--categories",
                "--out",
                str(out),
                "--out-csv",
                str(tmp_path / "out.csv"),
            ]
        )
        assert code == 0
        run = json.loads(out.read_text(encoding="utf-8"))
        assert run["input"]["chunks"] == 2
        assert "chunks of 2,000 rows (2 chunks)" in capsys.readouterr().out


class TestPeakRss:
    def test_peak_rss_sees_an_allocation(self):
        with PeakRss() as peak:
            block = bytearray(200 * 2**20)
            block[::4096] = b"x" * len(block[::4096])
        assert peak.peak_delta_bytes > 100 * 2**20
