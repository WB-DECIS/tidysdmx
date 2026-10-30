"""Scrub the real WB_IDS mapping workbook: ``python -m benchmarks.wb_ids.template``.

The mapping itself — series codes, their target codes, the component rules — is
public IDS metadata and is kept as is, so the mock pipeline maps exactly as the
real one does. Two kinds of change are made:

- Personal data is replaced: the INFO sheet's TTL name and email, any
  ``mailto:`` hyperlink, and the workbook's author properties. The custom
  document properties are dropped: the real workbook carries the World Bank
  tenant's sensitivity label and SharePoint metadata there.
- Typos that silently lose data in the real workbook are fixed, so the mock
  pipeline maps everything it is meant to (docs/datasets/wb-ids.md lists both):
  - A COMP_MAPPING ``SOURCE`` that matches a bronze CSV column only when case
    is ignored is respelled as the column. ``map_structures`` matches sources
    exactly, so the real ``Country_code`` (the CSV has ``Country_Code``) drops
    REF_AREA.
  - Codes in REP_MAPPING's ``S:``/``T:`` columns and in REQUESTED_IND_RAW lose
    surrounding whitespace. Three real series codes end in non-breaking
    spaces, so they never match the data.
"""

import argparse
import sys
from collections.abc import Sequence
from itertools import pairwise
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.packaging.custom import CustomPropertyList
from openpyxl.worksheet.worksheet import Worksheet

from benchmarks.wb_ids import TEMPLATE_XLSX
from benchmarks.wb_ids.profile import COLUMNS

PLACEHOLDERS: dict[str, str] = {
    "TTL NAME": "Mock Task Team Leader",
    "TTL EMAIL": "ttl@example.org",
}
AUTHOR: str = "tidysdmx benchmarks"
# Sheets holding codes, and which of their columns do: a header prefix, or ""
# for every column.
CODE_SHEETS: dict[str, tuple[str, ...]] = {
    "REP_MAPPING": ("S:", "T:"),
    "REQUESTED_IND_RAW": ("",),
}


def scrub(real: Path, out: Path) -> list[str]:
    """Write a copy of the real workbook with personal data replaced.

    Args:
        real: The real WB_IDS mapping workbook.
        out: Where to write the scrubbed copy.

    Returns:
        One line per change made, for the reader.

    Raises:
        ValueError: If the workbook lacks the INFO or COMP_MAPPING sheet.
    """
    workbook = load_workbook(real)
    for sheet in ("INFO", "COMP_MAPPING"):
        if sheet not in workbook.sheetnames:
            raise ValueError(f"{real} has no {sheet} sheet; expected a WB template")
    changes = _replace_info(workbook["INFO"])
    changes += _drop_mailto_links(workbook)
    changes += _align_sources(workbook["COMP_MAPPING"])
    for sheet in CODE_SHEETS:
        if sheet in workbook.sheetnames:
            changes += _strip_codes(workbook[sheet])
    workbook.properties.creator = AUTHOR
    workbook.properties.lastModifiedBy = AUTHOR
    changes.append(f"properties: creator and lastModifiedBy -> {AUTHOR!r}")
    changes += _drop_custom_properties(workbook)
    workbook.save(out)
    return changes


def _drop_custom_properties(workbook: Workbook) -> list[str]:
    """Remove every custom document property, sensitivity labels included."""
    names = workbook.custom_doc_props.names
    if not names:
        return []
    # An empty list is not written, so the copy has no docProps/custom.xml.
    workbook.custom_doc_props = CustomPropertyList()
    return [f"custom properties: {len(names)} removed ({', '.join(names)})"]


def _replace_info(sheet: Worksheet) -> list[str]:
    """Replace the value beside each INFO key listed in ``PLACEHOLDERS``."""
    changes: list[str] = []
    for row in sheet.iter_rows():
        for cell, neighbour in pairwise(row):
            key = str(cell.value).strip().upper() if cell.value is not None else ""
            if key in PLACEHOLDERS and neighbour.value is not None:
                neighbour.value = PLACEHOLDERS[key]
                changes.append(f"INFO!{neighbour.coordinate}: {key} -> placeholder")
    return changes


def _drop_mailto_links(workbook: Workbook) -> list[str]:
    """Remove every ``mailto:`` hyperlink, which would keep the real address."""
    changes: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                link = cell.hyperlink
                if link is not None and str(link.target or "").startswith("mailto:"):
                    cell.hyperlink = None
                    changes.append(f"{sheet.title}!{cell.coordinate}: mailto removed")
    return changes


def _strip_codes(sheet: Worksheet) -> list[str]:
    """Strip surrounding whitespace, non-breaking spaces included, from codes."""
    prefixes = CODE_SHEETS[sheet.title]
    header = [str(cell.value or "") for cell in sheet[1]]
    columns = [i for i, name in enumerate(header) if name.startswith(prefixes) and name]
    changes: list[str] = []
    for row in sheet.iter_rows(min_row=2):
        for i in columns:
            cell = row[i]
            if isinstance(cell.value, str) and cell.value != cell.value.strip():
                changes.append(
                    f"{sheet.title}!{cell.coordinate}: {cell.value!r} stripped"
                )
                cell.value = cell.value.strip()
    return changes


def _align_sources(sheet: Worksheet) -> list[str]:
    """Respell each SOURCE that matches a CSV column only when case is ignored."""
    header = [cell.value for cell in sheet[1]]
    if "SOURCE" not in header:
        raise ValueError("COMP_MAPPING has no SOURCE column")
    column = header.index("SOURCE") + 1
    by_lower = {name.lower(): name for name in COLUMNS}
    changes: list[str] = []
    for (cell,) in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
        value = cell.value
        if not isinstance(value, str):
            continue
        spelled = by_lower.get(value.lower())
        if spelled is not None and spelled != value:
            cell.value = spelled
            changes.append(f"COMP_MAPPING!{cell.coordinate}: {value!r} -> {spelled!r}")
    return changes


def main(argv: Sequence[str] | None = None) -> int:
    """Scrub ``--real`` into ``--out``.

    Args:
        argv: Command-line arguments; ``sys.argv[1:]`` when ``None``.

    Returns:
        0 on success.
    """
    parser = argparse.ArgumentParser(
        prog="python -m benchmarks.wb_ids.template",
        description="Scrub the real WB_IDS mapping workbook into the mock one.",
    )
    parser.add_argument("--real", type=Path, required=True)
    parser.add_argument("--out", type=Path, default=TEMPLATE_XLSX)
    args = parser.parse_args(argv)
    for change in scrub(args.real, args.out):
        sys.stderr.write(change + "\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
